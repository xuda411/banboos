from io import BytesIO

from fastapi.testclient import TestClient
from openpyxl import load_workbook

from apps.api import main
from apps.worker.main import run_once


def test_portfolio_api_worker_and_export_snapshot():
    client = TestClient(main.app)
    payload = {'kind': 'portfolio-optimization', 'parameters': {
        'objective': 'max_npv', 'budget_limit_wan': 1000,
        'projects': [{'name': '=手工假设', 'capacity_mwh': 10,
                      'unit_investment_yuan_wh': 1, 'annual_revenue_wan': 300}]}}
    response = client.post('/api/v1/runs', json=payload)
    assert response.status_code == 202
    run_id = response.json()['run_id']
    assert client.get(f'/api/v1/runs/{run_id}/export').status_code == 409
    assert run_once(main.run_registry)
    completed = client.get(f'/api/v1/runs/{run_id}').json()
    assert completed['status'] == 'succeeded'
    assert completed['result']['selected_projects'][0]['annual_revenue_wan'] == 300
    exported = client.get(f'/api/v1/runs/{run_id}/export')
    assert exported.status_code == 200
    book = load_workbook(BytesIO(exported.content))
    assert book.sheetnames == ['导出说明', '任务参数', '结果摘要', '全部候选项目', '选中项目', '组合复核']
    assert book['选中项目']['A5'].value == '=手工假设'
    assert book['选中项目']['A5'].data_type == 's'
    assert book['选中项目']['D5'].value == 300
    assert book['选中项目']['E5'].value == 1000
    assert book['全部候选项目'].freeze_panes == 'B5'
    assert book['组合复核']['B10'].value == 1
    assert book['组合复核']['B14'].value == 'optimal'
    assert book['导出说明']['B5'].value == run_id
    payload['parameters']['budget_limit_wan'] = None
    assert client.post('/api/v1/runs', json=payload).status_code == 422


def test_portfolio_no_selection_exports_reason():
    client = TestClient(main.app)
    submitted = client.post('/api/v1/runs', json={'kind': 'portfolio-optimization', 'parameters': {
        'objective': 'min_investment', 'revenue_target_wan': 500,
        'projects': [{'name': '甲', 'capacity_mwh': 10,
                      'unit_investment_yuan_wh': 1, 'annual_revenue_wan': 0}]}})
    assert submitted.status_code == 202
    run_id = submitted.json()['run_id']
    assert run_once(main.run_registry)
    result = client.get(f'/api/v1/runs/{run_id}').json()
    assert result['result']['outcome'] == 'infeasible'
    report = client.get(f'/api/v1/runs/{run_id}/export')
    assert report.status_code == 200
    book = load_workbook(BytesIO(report.content))
    assert book['选中项目'].max_row == 4
    assert any(row[0] == 'outcome' and row[1] == 'infeasible'
               for row in book['结果摘要'].iter_rows(min_row=5, values_only=True))
