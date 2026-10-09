from zipfile import ZipFile

import pytest
from openpyxl import load_workbook

from packages.application.financial_export import export_financial_xlsx
from packages.application.financial_template_xlsm import (
    export_financial_xlsm,
    financial_template_path,
    template_input_values,
    template_sheet_values,
)
from packages.application.financial_xlsm_review import review_native_cached_values
from packages.contracts.financial import FinancialTaskParameters
from packages.domain.financial_model import calculate_financials


def test_financial_export_has_typed_sheets_and_frozen_headers(tmp_path):
    result = calculate_financials(FinancialTaskParameters(
        power_mw=100, capacity_mwh=200, annual_revenue_yuan=8_000_000,
    ).financial())
    result["source_run_id"] = "demo-run"
    result["input_parameters"] = {"power_mw": 100, "capacity_mwh": 200, "annual_revenue_yuan": 8_000_000}
    destination = export_financial_xlsx(result, tmp_path / "financial.xlsx")
    workbook = load_workbook(destination, data_only=False)
    assert workbook.sheetnames == ["导出说明", "项目概览", "测算参数", "年度现金流", "融资明细", "全周期现金流", "财务指标", "财务模板", "模板复核", "重算复核", "模板映射"]
    assert workbook["导出说明"]["B5"].value == "banboos-export-2026-09-21"
    assert workbook["测算参数"]["B5"].value == 100
    assert workbook["年度现金流"].freeze_panes == "B5"
    assert workbook["年度现金流"]["C5"].value == result["yearly"][0]["energy_revenue_yuan"]
    assert workbook["年度现金流"]["U4"].value == "实际应缴增值税（元）"
    assert workbook["年度现金流"]["AB4"].value == "项目税前现金流（元）"
    assert workbook["项目概览"]["B12"].value == result["total_investment_yuan"]
    statement = workbook["财务指标"]
    assert statement.freeze_panes == "D5"
    assert statement["D31"].value == "='全周期现金流'!B5"
    assert statement["E31"].value == "='年度现金流'!AB5"
    assert statement["D39"].value.startswith("=IRR(D31:")
    assert statement["D42"].value.startswith("=NPV('测算参数'!")
    parity = workbook["财务模板"]
    assert parity.freeze_panes == "D3"
    assert parity["D2"].value == 0
    assert parity["E2"].value == 1
    assert parity["C27"].value == "收益合计（万元）"
    assert parity["E27"].value.startswith("=SUM(E10,E14,E18,E22,E26)")
    assert parity["D76"].value == "='财务指标'!D32/10000"
    assert parity["D78"].value.startswith("=IRR(D76:")
    assert "E40" not in parity["E34"].value
    assert parity["E56"].value == "=SUM(E57:E62)"
    assert "A6:A27" in {str(range_ref) for range_ref in parity.merged_cells.ranges}
    assert parity.column_dimensions["A"].width == 18
    assert parity.column_dimensions["D"].width == 12
    assert parity["A2"].fill.fgColor.rgb == "00B4C6E7"
    assert parity["D78"].font.color.rgb == "00C00000"
    audit = workbook["模板复核"]
    assert audit["B5"].value == "PASS"
    assert audit["B8"].value == "PASS"
    assert audit["B11"].value == "PASS"
    assert audit["B13"].value == "WARN"
    recalc = workbook["重算复核"]
    assert recalc["B5"].value == "='财务模板'!D78"
    assert recalc["C5"].value == result["full_irr"]
    assert recalc["D5"].value == "=B5-C5"
    assert recalc["B15"].value == "PASS"
    assert recalc["B17"].value == "WARN"


def test_formula_audit_uses_task_inputs_and_preserves_snapshot(tmp_path):
    parameters = FinancialTaskParameters(power_mw=30, capacity_mwh=120,
        annual_revenue_yuan=20_000_000, loan_ratio=.7, replace_year=3,
        replace_capex_yuan=8_000_000)
    result = calculate_financials(parameters.financial())
    result["input_parameters"] = parameters.model_dump(exclude_none=True)
    path = export_financial_xlsx(result, tmp_path / "audit.xlsx")
    workbook = load_workbook(path)
    audit = workbook["公式复核"]
    assert audit["C5"].data_type == "f"
    assert "测算参数" in audit["C5"].value
    assert audit["M5"].value == "=C5-D5-F5-G5-H5-K5"
    assert audit["O29"].value == "=M29-'年度现金流'!O29"
    assert workbook["年度现金流"]["O5"].value == result["yearly"][0]["equity_cashflow_yuan"]
    assert "年限变更" in audit["A2"].value
    result["model_version"] = "banboos-financial-1.0.0"
    old = load_workbook(export_financial_xlsx(result, tmp_path / "old.xlsx"))
    assert "公式复核" not in old.sheetnames


def test_formula_audit_includes_template_cost_and_tax_drivers(tmp_path):
    parameters = FinancialTaskParameters(
        power_mw=100, capacity_mwh=200, annual_revenue_yuan=12_000_000,
        insurance_rate=0.002, fixed_operation_cost_yuan=100_000,
        revenue_share_threshold_yuan=5_000_000, revenue_share_rate=0.3,
        vat_rate=0.13, stamp_tax_rate=0.0005,
        eol_method="calendar_cycle_min",
    )
    result = calculate_financials(parameters.financial())
    result["input_parameters"] = parameters.model_dump(exclude_none=True)
    workbook = load_workbook(export_financial_xlsx(result, tmp_path / "template-parity.xlsx"))
    audit = workbook["公式复核"]
    assert "MAX(0" in audit["J5"].value
    assert "B5" in audit["D5"].value
    assert audit["B5"].value.startswith("=IF(")
    assert "Q5" in audit["P5"].value
    assert audit["S5"].value.startswith("=(C5-")


def test_phase_schedule_is_serialized_and_rebuilt_in_formula_audit(tmp_path):
    parameters = FinancialTaskParameters(
        power_mw=10, capacity_mwh=20, annual_revenue_yuan=5_000_000,
        capacity_fee_yuan=1_000_000,
        revenue_phases={"capacity_fee_yuan": {"start_year": 2, "end_year": 4}},
    )
    result = calculate_financials(parameters.financial())
    result["input_parameters"] = parameters.model_dump(exclude_none=True)
    workbook = load_workbook(export_financial_xlsx(result, tmp_path / "phase.xlsx"))
    parameter_values = [workbook["测算参数"].cell(row, 2).value
                        for row in range(5, workbook["测算参数"].max_row + 1)]
    assert any(isinstance(value, str) and "capacity_fee_yuan" in value for value in parameter_values)
    assert "阶段收益" in workbook.sheetnames
    assert "公式复核" in workbook.sheetnames
    assert "SUMIFS" in workbook["公式复核"]["C5"].value
    assert "阶段收益" in workbook["公式复核"]["C5"].value


def test_template_xlsm_mapping_uses_native_units():
    parameters = FinancialTaskParameters(
        power_mw=100, capacity_mwh=200, annual_revenue_yuan=8_000_000,
        capacity_fee_yuan=1_000_000, loan_ratio=.7,
        replace_year=3, replace_capex_yuan=8_000_000,
    )
    values = template_input_values({"input_parameters": parameters.model_dump(exclude_none=True)})
    assert values["D3"] == 100
    assert values["D4"] == 200
    assert values["D40"] == .04
    assert values["D44"] == .7
    assert values["D50"] == 100
    assert values["D51"] == 800
    assert values["D38"] == "是"


def test_template_xlsm_maps_desktop_eol_curves():
    parameters = FinancialTaskParameters(
        power_mw=100, capacity_mwh=200, annual_revenue_yuan=8_000_000,
        eol_method="desktop_template", calendar_eol_table=[1.0, 0.995, 0.99],
    )
    mappings = template_sheet_values({"input_parameters": parameters.model_dump(exclude_none=True),
                                      "yearly": [{}] * 25})
    assert mappings["EOL"]["C3"] == 1.0
    assert mappings["EOL"]["D3"] == 0.995
    assert mappings["EOL"]["E4"] == pytest.approx(0.9725)
    assert mappings["EOL"]["D9"] == pytest.approx(0.99)
    assert mappings["电量类 "]["D24"] == pytest.approx(776)


def test_template_xlsm_mapping_accepts_native_mode():
    parameters = FinancialTaskParameters(
        power_mw=60, capacity_mwh=120, annual_revenue_yuan=16_919_681.739130434,
        eol_method="native_xlsm",
    )
    mappings = template_sheet_values({"input_parameters": parameters.model_dump(exclude_none=True),
                                      "yearly": [{}] * 25})
    assert mappings["电量类 "]["D24"] == pytest.approx(16_919_681.739130434 / 10000 * .97)
    assert mappings["EOL"]["D9"] == pytest.approx(.99)


def test_template_xlsm_copies_source_and_records_native_review(tmp_path):
    source = financial_template_path()
    if source is None:
        pytest.skip("开发机未提供 1.6.6 原版 XLSM 模板")
    parameters = FinancialTaskParameters(power_mw=100, capacity_mwh=200,
                                         annual_revenue_yuan=8_000_000)
    result = calculate_financials(parameters.financial())
    result["input_parameters"] = parameters.model_dump(exclude_none=True)
    result["run_id"] = "xlsm-test"
    destination = export_financial_xlsm(result, tmp_path / "financial.xlsm", source)
    workbook = load_workbook(destination, data_only=False, keep_vba=True)
    assert "Banboos2.0快照" in workbook.sheetnames
    assert "原版对账" in workbook.sheetnames
    assert "输入映射记录" in workbook.sheetnames
    assert workbook["参数设定 "]["D3"].value == 100
    assert workbook["参数设定 "]["D4"].value == 200
    assert workbook["容量类"]["D25"].value == 0
    assert workbook["电量类 "]["D24"].value == 800
    assert workbook["辅助服务类"]["D37"].value == 0
    with ZipFile(destination) as archive:
        assert not any(name.lower().endswith("vbaproject.bin") for name in archive.namelist())
    cached = review_native_cached_values(destination, result)
    assert cached["engine_executed"] is False
    assert cached["formula_equivalence_verified"] is False
    assert cached["status"] in {"PENDING_RECALCULATION", "DIFFERENCES", "CACHED_VALUES_MATCH"}
