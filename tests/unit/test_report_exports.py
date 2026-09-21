from io import BytesIO

from openpyxl import load_workbook

from packages.application.report_export import (
    export_operations_report_xlsx,
    export_portfolio_candidates_xlsx,
    export_price_aggregate_xlsx,
)
from packages.contracts.operations_report import (
    OperationsReport,
    OperationsReportPeriod,
    OperationsReportProvince,
)
from packages.contracts.portfolio_candidates import PortfolioCandidate
from packages.contracts.portfolio_candidates_result import PortfolioCandidatesResult
from packages.contracts.readonly import PriceAggregatePeriod, PriceAggregateResult


def _sheets(content):
    return load_workbook(BytesIO(content), data_only=False)


def test_price_aggregate_export_keeps_period_and_coverage_audit():
    result = PriceAggregateResult(
        node_id=7, market="实时", start_date="2026-01-01", end_date="2026-01-31",
        duration_hours=2, valid_days=3, available_days=4, excluded_records=1,
        multiple_source_days=1, multiple_source_dates=["2026-01-02"],
        missing_dates=["2026-01-04"],
        monthly=[PriceAggregatePeriod(period="2026-01", valid_days=3,
                  average_price_yuan_per_mwh=300, charge_price_yuan_per_mwh=100,
                  discharge_price_yuan_per_mwh=500, spread_yuan_per_mwh=400,
                  min_price_yuan_per_mwh=10, max_price_yuan_per_mwh=900,
                  multiple_source_days=1)], annual=[], source_mode="staging-readonly")
    workbook = _sheets(export_price_aggregate_xlsx(result))
    assert workbook.sheetnames == ["导出说明", "月度统计", "年度统计", "覆盖审计"]
    assert workbook["月度统计"]["A4"].value == "周期类型"
    assert workbook["月度统计"]["B5"].value == "2026-01"
    assert workbook["覆盖审计"]["B5"].value == "2026-01-04"
    assert any(row[0].value == "导出格式版本" and row[1].value == "banboos-export-2026-09-21"
               for row in workbook["导出说明"].iter_rows(min_row=5))


def test_operations_export_has_province_and_monthly_views():
    result = OperationsReport(
        market="实时", start_date="2026-01-01", end_date="2026-01-31", duration_hours=2,
        node_count=2, valid_nodes=1, total_valid_days=5, total_data_points=480,
        provinces=[OperationsReportProvince(province="湖北", node_count=2, valid_nodes=1,
                    valid_days=5, data_points=480)],
        monthly=[OperationsReportPeriod(period="2026-01", node_count=1, valid_days=5,
                 average_spread_yuan_per_mwh=420)], source_mode="staging-readonly")
    workbook = _sheets(export_operations_report_xlsx(result))
    assert workbook.sheetnames == ["导出说明", "省级汇总", "月度价差"]
    assert workbook["省级汇总"]["A5"].value == "湖北"
    assert workbook["月度价差"]["D5"].value == 420


def test_portfolio_export_exposes_snapshot_and_scale_check():
    result = PortfolioCandidatesResult(
        market="实时", power_mw=100, capacity_mwh=200, duration_hours=2,
        round_trip_efficiency=.92, start_date="2026-01-01", end_date="2026-01-31",
        snapshot_id="a" * 64, algorithm_version="portfolio-candidates-v1",
        candidates=[PortfolioCandidate(name="湖北·实时", node_id=7, province="湖北", market="实时",
            capacity_mwh=200, unit_investment_yuan_wh=1.2, annual_revenue_wan=800,
            spread_yuan_per_mwh=400, valid_days=31, source_mode="staging-readonly")])
    workbook = _sheets(export_portfolio_candidates_xlsx(result))
    assert workbook.sheetnames == ["导出说明", "候选项目", "规模校验"]
    assert workbook["候选项目"]["A5"].value == "湖北·实时"
    assert workbook["规模校验"]["C7"].value == "PASS"
