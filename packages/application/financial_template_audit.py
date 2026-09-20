"""Read-only audit of the server workbook against the supplied xlsm layout."""
from __future__ import annotations

import re

from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

THIN = Side(style="thin", color="C9D8D2")
GREEN = "0B6B53"
TEXT = "173B35"
PASS = "E2F0D9"
WARN = "FFF2CC"
FAIL = "FCE4D6"


def write_template_audit(workbook) -> None:
    """Add a terminal, human-readable parity report without reading or changing the xlsm."""
    sheet = workbook.create_sheet("模板复核")
    sheet.cell(1, 1, "原始 xlsm 模板对照复核")
    sheet.cell(1, 1).font = Font(name="Microsoft YaHei", size=16, bold=True, color=TEXT)
    sheet.merge_cells("A1:E1")
    sheet.cell(2, 1, "本页只检查服务器导出结构、公式和样式；原始 xlsm 保持只读，宏/数据表/数组公式差异单独标明。")
    sheet.merge_cells("A2:E2")
    headers = ["检查项", "状态", "服务器导出", "原始模板基线", "说明"]
    for column, value in enumerate(headers, 1):
        cell = sheet.cell(4, column, value)
        cell.fill = PatternFill("solid", fgColor=GREEN)
        cell.font = Font(name="Microsoft YaHei", bold=True, color="FFFFFF")
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = Border(bottom=THIN)

    target = workbook["财务模板"] if "财务模板" in workbook.sheetnames else None
    checks: list[tuple[str, str, str, str, str]] = []

    def add(name: str, ok: bool, actual: str, baseline: str, note: str, *, warning: bool = False) -> None:
        checks.append((name, "WARN" if warning else ("PASS" if ok else "FAIL"), actual, baseline, note))

    add("工作表存在", target is not None,
        "存在" if target else "缺少", "财务指标结构页", "服务器导出使用独立财务模板页。")
    if target is None:
        _write_rows(sheet, checks)
        return

    add("行列范围", target.max_row == 81 and target.max_column == 30,
        f"{target.max_row} 行 × {target.max_column} 列", "81 行 × 30 列",
        "对应原始财务指标页的 0 年、25 个运营年度和合计列。")
    expected_merges = {"A1:AD1", "A3:C3", "A4:A5", "A6:A27", "A29:A32",
                       "A33:A40", "A41:A44", "A45:A53", "A54:C54", "A55:A67",
                       "A68:C68", "A69:A80", "A81:F81"}
    actual_merges = {str(value) for value in target.merged_cells.ranges}
    add("合并区域", expected_merges.issubset(actual_merges),
        f"{len(actual_merges)} 个，关键区域 {len(expected_merges & actual_merges)} 个",
        f"至少 {len(expected_merges)} 个关键区域", "分类、标题和说明区域采用模板同样的合并关系。")
    labels = {
        "A3": "1、项目充放电数据与收支明细", "C4": "电池容量衰减 EOL",
        "C27": "收益合计（万元）", "C40": "总运营费用（万元）",
        "C45": "应缴纳增值税（万元）", "C63": "现金净流入（万元）",
        "C76": "现金净流入（万元）",
    }
    labels_ok = all(target[cell].value == value for cell, value in labels.items())
    add("关键标签", labels_ok, "全部匹配" if labels_ok else "存在差异",
        "7 个关键标签", "用于防止模板行移动后仍被误认为已对齐。")
    key_formulas = {
        "D65": "IRR", "D66": "NPV", "D78": "IRR", "D79": "NPV",
        "D63": "财务指标", "D76": "财务指标",
    }
    formula_ok = all(isinstance(target[cell].value, str)
                     and target[cell].value.startswith("=")
                     and marker in target[cell].value
                     for cell, marker in key_formulas.items())
    formula_count = sum(1 for row in target.iter_rows() for cell in row
                        if isinstance(cell.value, str) and cell.value.startswith("="))
    add("关键公式", formula_ok, f"关键公式通过；共 {formula_count} 个公式",
        "IRR/NPV/现金流链接公式", "金额公式链接服务器计算页，不复制原始宏。")
    self_refs = _self_references(target)
    add("公式自引用", not self_refs, str(len(self_refs)), "0",
        "检测同一单元格直接引用自身，避免新增模板层形成循环。")
    widths_ok = target.column_dimensions["A"].width == 18 and target.column_dimensions["C"].width == 28
    add("列宽", widths_ok, f"A={target.column_dimensions['A'].width}, C={target.column_dimensions['C'].width}",
        "A≈18, C≈28", "模板页采用固定列宽，避免公式文本撑开年度列。")
    color_ok = (target["A2"].fill.fgColor.rgb == "00B4C6E7"
                and target["A3"].fill.fgColor.rgb == "00FFF2CC"
                and target["D78"].font.color.rgb == "00C00000")
    add("颜色与强调", color_ok, "浅蓝表头/浅黄分组/红色 IRR" if color_ok else "存在差异",
        "原始模板主要视觉规则", "当前只复制稳定的 RGB 视觉规则，不依赖主题色。")
    add("宏与数据表", False, "未写入 XLSX", "原始 xlsm 含 VBA 与 DataTableFormula",
        "这是有意保留的能力边界，不能用静态 XLSX 冒充宏或数据表。", warning=True)
    add("数组公式", False, "未写入 XLSX", "原始模板含回收期数组公式",
        "服务器版使用可复核的静态回收期结果，原始数组公式仍需 Excel/WPS 专项验证。", warning=True)
    _write_rows(sheet, checks)


def _self_references(sheet) -> list[str]:
    reference = re.compile(r"(?<![A-Z0-9_])\$?([A-Z]{1,3})\$?(\d+)")
    issues: list[str] = []
    for row in sheet.iter_rows():
        for cell in row:
            if not isinstance(cell.value, str) or not cell.value.startswith("="):
                continue
            refs = {f"{match.group(1)}{match.group(2)}" for match in reference.finditer(cell.value)}
            if cell.coordinate in refs:
                issues.append(cell.coordinate)
    return issues


def _write_rows(sheet, checks: list[tuple[str, str, str, str, str]]) -> None:
    for row, values in enumerate(checks, 5):
        for column, value in enumerate(values, 1):
            cell = sheet.cell(row, column, value)
            cell.font = Font(name="Microsoft YaHei", color=TEXT,
                             bold=column == 2)
            cell.fill = PatternFill("solid", fgColor={"PASS": PASS, "WARN": WARN,
                                                        "FAIL": FAIL}.get(values[1], "FFFFFF"))
            cell.border = Border(bottom=THIN)
            cell.alignment = Alignment(vertical="center", wrap_text=True)
    sheet.auto_filter.ref = f"A4:E{sheet.max_row}"
    sheet.freeze_panes = "A5"
    for column, width in {"A": 20, "B": 10, "C": 28, "D": 34, "E": 58}.items():
        sheet.column_dimensions[column].width = width
