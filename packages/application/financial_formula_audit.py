"""Excel formulas independently rebuild cashflows from the exported task inputs."""
from openpyxl.formatting.rule import CellIsRule
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

from packages.domain.financial_model import MODEL_VERSION, FinancialParameters


def write_formula_audit(workbook, result, title, header):
    """Keep historic snapshots intact; expose a recalculable audit for the corrected model."""
    if result.get("model_version") != MODEL_VERSION or not result.get("input_parameters"):
        return
    required = set(FinancialParameters.__dataclass_fields__) - {"replace_year"}
    if not required.issubset(result["input_parameters"]):
        return  # Old or incomplete inputs cannot support a trustworthy independent reconstruction.
    parameters = workbook["测算参数"]
    refs = {parameters.cell(row, 4).value: f"'测算参数'!$B${row}"
            for row in range(5, parameters.max_row + 1)}
    sheet = workbook.create_sheet("公式复核")
    title(sheet, "逐年现金流公式复核", 22)
    sheet.cell(2, 1, "本页由测算参数重算；项目概览及其余明细保留任务快照。差额不为零时请在软件重新测算。年限变更须重新导出。")
    headers = ["年度", "EOL", "总收入（元）", "运维成本（元）", "折旧（元）",
               "换电池投资（元）", "贷款利息（元）", "偿还本金（元）", "剩余贷款（元）",
               "项目所得税（元）", "资本金所得税（元）", "项目现金流（元）",
               "资本金现金流（元）", "项目快照差额（元）", "资本金快照差额（元）"]
    headers += ["实际应缴增值税（元）", "进项税抵扣期初（元）", "进项税抵扣期末（元）",
                "项目税前现金流（元）", "资本金税前现金流（元）",
                "项目税前快照差额（元）", "资本金税前快照差额（元）"]
    header(sheet, 4, headers)
    p = refs.__getitem__
    initial = f"({p('capacity_mwh')}*1000000*{p('capex_yuan_per_wh')})"
    loan = f"({initial}*{p('loan_ratio')})"
    total = f"({initial}+{loan}*{p('construction_loan_rate')}*{p('construction_years')})"
    years = p("operation_years")
    replace_year, replace_capex = p("replace_year"), p("replace_capex_yuan")
    inputs = result["input_parameters"]
    extended = any(float(inputs.get(key, default)) != default for key, default in {
        "insurance_rate": 0.0, "fixed_operation_cost_yuan": 0.0,
        "revenue_share_threshold_yuan": 0.0, "revenue_share_rate": 0.0,
        "other_operating_cost_yuan": 0.0, "vat_rate": 0.0,
        "stamp_tax_rate": 0.0, "calendar_eol_decline": 0.015,
        "annual_cycles": 350.0, "cycle_life_cycles": 8000.0,
        "input_vat_rate_equipment": 0.13, "input_vat_rate_other": 0.09,
        "equipment_investment_share": 1.0, "input_vat_credit_ratio": 1.0,
    }.items()) or inputs.get("eol_method", "linear") != "linear"
    for row in range(5, len(result["yearly"]) + 5):
        year = row - 4
        previous = loan if year == 1 else f"I{row - 1}"
        gross_revenue = f"({p('annual_revenue_yuan')}+{p('capacity_fee_yuan')}+{p('subsidy_yuan')}+{p('capacity_lease_yuan')}+{p('primary_frequency_yuan')}+{p('secondary_frequency_yuan')})*B{row}"
        output_vat = f"({gross_revenue}/(1+{p('vat_rate')})*{p('vat_rate')})"
        net_revenue = f"({gross_revenue}-{output_vat})"
        vat_surcharge = f"({output_vat}*{p('vat_surcharge_rate')})"
        stamp_tax = f"({net_revenue}*{p('stamp_tax_rate')})"
        revenue_share = f"MAX(0,({net_revenue}-{p('revenue_share_threshold_yuan')})*{p('revenue_share_rate')})"
        insurance = f"({initial}*{p('insurance_rate')})"
        initial_input_credit = (
            f"({initial}*{p('equipment_investment_share')}/(1+{p('input_vat_rate_equipment')})*"
            f"{p('input_vat_rate_equipment')}+{initial}*(1-{p('equipment_investment_share')})/"
            f"(1+{p('input_vat_rate_other')})*{p('input_vat_rate_other')})*"
            f"{p('input_vat_credit_ratio')}"
        )
        input_vat_new = (f"(F{row}/(1+{p('input_vat_rate_equipment')})*"
                         f"{p('input_vat_rate_equipment')}*{p('input_vat_credit_ratio')})")
        input_credit_opening = initial_input_credit if year == 1 else f"R{row - 1}"
        actual_vat = f"({output_vat}-MIN({output_vat},Q{row}+{input_vat_new}))"
        vat_surcharge = f"({actual_vat}*{p('vat_surcharge_rate')})"
        project_pre_tax = f"(C{row}-D{row}-F{row}-P{row}-{vat_surcharge}-{stamp_tax})"
        equity_pre_tax = f"(C{row}-D{row}-F{row}-G{row}-H{row}-P{row}-{vat_surcharge}-{stamp_tax})"
        formulas = {
            "A": str(year),
            "B": f"IF({p('eol_method')}=\"linear\",{p('first_year_eol')}+({p('final_eol')}-{p('first_year_eol')})*(A{row}-1)/MAX(1,{years}-1),MIN(MAX({p('final_eol')},1-{p('calendar_eol_decline')}*A{row}),MAX({p('final_eol')},1-{p('annual_cycles')}/{p('cycle_life_cycles')}*A{row})))",
            "C": gross_revenue,
            "D": f"{total}*{p('om_rate')}*(1+{p('om_growth')})^(A{row}-1)+{p('land_rent_yuan')}+{insurance}+{p('fixed_operation_cost_yuan')}+{p('other_operating_cost_yuan')}+{revenue_share}",
            "E": f"{initial}*(1-{p('residual_rate')})/{years}+IF(AND({replace_year}>0,A{row}>={replace_year}),{replace_capex}*(1-{p('residual_rate')})/MAX(1,{years}-{replace_year}+1),0)",
            "F": f"IF(A{row}={replace_year},{replace_capex},0)",
            "G": f"IF(A{row}<={p('loan_years')},{previous}*{p('loan_rate')},0)",
            "H": f"IF(A{row}={years},{previous},IF(A{row}<={p('loan_years')},{loan}/{p('loan_years')},0))",
            "I": f"MAX(0,{previous}-H{row})",
            "J": f"MAX(0,({net_revenue}-D{row}-E{row}-{vat_surcharge}-{stamp_tax})*{p('income_tax_rate')})",
            "K": f"MAX(0,({net_revenue}-D{row}-E{row}-G{row}-{vat_surcharge}-{stamp_tax})*{p('income_tax_rate')})",
            "L": f"S{row}-J{row}",
            "M": f"T{row}-K{row}",
            "N": f"L{row}-'年度现金流'!N{row}",
            "O": f"M{row}-'年度现金流'!O{row}",
            "P": actual_vat,
            "Q": input_credit_opening,
            "R": f"MAX(0,Q{row}+{input_vat_new}-{output_vat})",
            "S": project_pre_tax,
            "T": equity_pre_tax,
            "U": f"S{row}-'年度现金流'!AB{row}",
            "V": f"T{row}-'年度现金流'!AC{row}",
        }
        if not extended:
            formulas.update({
                "B": f"{p('first_year_eol')}+({p('final_eol')}-{p('first_year_eol')})*(A{row}-1)/MAX(1,{years}-1)",
                "D": f"{total}*{p('om_rate')}*(1+{p('om_growth')})^(A{row}-1)",
                "J": f"MAX(0,(C{row}-D{row}-E{row})*{p('income_tax_rate')})",
                "K": f"MAX(0,(C{row}-D{row}-E{row}-G{row})*{p('income_tax_rate')})",
                "L": f"C{row}-D{row}-J{row}-F{row}",
                "M": f"C{row}-D{row}-F{row}-G{row}-H{row}-K{row}",
            })
        for column, formula in formulas.items():
            cell = sheet[f"{column}{row}"]
            cell.value = year if column == "A" else f"={formula}"
            cell.font = Font(name="Microsoft YaHei", size=10, color="0B6B53" if "'" in formula else "173B35")
            cell.number_format = "0.00%" if column == "B" else "#,##0.00;[Red](#,##0.00);0.00"
        sheet[f"A{row}"].number_format = '0" 年"'
    sheet.auto_filter.ref = f"A4:{get_column_letter(len(headers))}{sheet.max_row}"
    sheet.conditional_formatting.add(f"N5:O{sheet.max_row}", CellIsRule(
        operator="notBetween", formula=["-0.01", "0.01"],
        fill=PatternFill("solid", fgColor="FCE4D6")))
    sheet.conditional_formatting.add(f"U5:V{sheet.max_row}", CellIsRule(
        operator="notBetween", formula=["-0.01", "0.01"],
        fill=PatternFill("solid", fgColor="FCE4D6")))
