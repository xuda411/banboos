"""Independent native-template checks. No Excel formula execution is claimed."""
from __future__ import annotations

from copy import copy
from hashlib import sha256
from math import isclose, isfinite
from pathlib import Path
from zipfile import ZipFile

from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

NATIVE_DIFFERENCES = [
    ("EOL!D3:AG5、财务指标!E5", "原表按手工日历/循环曲线取小，并用相邻年份平均值进入年度收益；desktop_template 已写入同一曲线，但年度插值仍需逐年验收。"),
    ("财务指标!E28、E50", "原表不含税收入固定除以 1.13，印花税固定 0.0005；服务器使用输入税率。"),
    ("财务指标!E49", "原表附加税按应计增值税的 12%；服务器按抵扣后的实缴税额及输入附加率。"),
    ("财务指标!E29、E32", "原表折旧含建设利息、末年回收残值；服务器折旧和残值现金流规则不同。"),
    ("财务指标!G34、E37", "原表运维费线性增长、按含税收入分成；服务器复利增长、按不含税收入分成。"),
    ("财务指标!D46、E46:E47", "进项税抵扣基数、抵扣比例及换电池当年抵扣顺序仍有差异。"),
    ("财务指标!E70、E73", "原表全投资现金流合计未含附加税/印花税，所得税引用含贷款利息的利润。"),
    ("财务指标!D65、D78", "原表标为税前 IRR，但引用扣过所得税的现金流；不能按标签认定税前等价。"),
    ("原生数据表/数组公式", "本程序不运行 Excel/WPS 引擎；缓存值对比与结构保留均不能证明公式完全等价。"),
]


def _difference_group(sheet: str, cell: str) -> str:
    """Classify a native-template difference for the next repair batch."""
    if sheet == "财务指标":
        row = int("".join(character for character in cell if character.isdigit()) or 0)
        if row in {4, 5, 8, 9, 12, 13, 14}:
            return "eol_and_revenue"
        if row in {29, 30, 31, 32}:
            return "depreciation_and_residual"
        if row in {45, 46, 47, 48, 49, 50, 51, 52, 53}:
            return "tax_and_deduction"
        if row in {63, 64, 65, 66, 76, 77, 78, 79, 80}:
            return "cashflow_and_return"
    return "other"


def comparison_cases(result: dict) -> list[dict]:
    cases = []

    def add(label, sheet, cell, expected, unit="万元"):
        cases.append({"label": label, "sheet": sheet, "cell": cell,
                      "expected": expected, "unit": unit,
                      "tolerance": 1e-7 if unit == "比例" else 1e-6})

    p = result["input_parameters"]
    add("初始投资", "参数设定 ", "D22", result["initial_investment_yuan"] / 10000)
    add("建设利息", "参数设定 ", "D18", result["construction_interest_yuan"] / 10000)
    add("总投资", "参数设定 ", "D24", result["total_investment_yuan"] / 10000)
    add("贷款本金", "参数设定 ", "D46", result["initial_investment_yuan"] * p["loan_ratio"] / 10000)
    for label, cell, key, unit in [
        ("全投资 IRR（原表税前标签存疑）", "D78", "full_irr", "比例"),
        ("资本金 IRR（原表税前标签存疑）", "D65", "equity_irr", "比例"),
        ("全投资 NPV", "D79", "full_npv_yuan", "万元"),
        ("资本金 NPV", "D66", "equity_npv_yuan", "万元"),
    ]:
        value = result.get(key)
        add(label, "财务指标", cell, value if unit == "比例" or value is None else value / 10000, unit)
    for cell, key in [("D76", "cashflows_yuan"), ("D63", "equity_cashflows_yuan")]:
        add("建设期现金流", "财务指标", cell, result[key][0] / 10000)
    fields = {5: ("EOL", "eol"), 6: ("容量租赁", "capacity_lease_yuan"),
              10: ("容量电费", "capacity_fee_yuan"), 14: ("电能量", "energy_revenue_yuan"),
              18: ("补贴", "subsidy_yuan"), 22: ("一次调频", "primary_frequency_yuan"),
              26: ("二次调频", "secondary_frequency_yuan"), 27: ("总收入", "revenue_yuan"),
              29: ("折旧", "depreciation_yuan"), 48: ("实缴增值税", "actual_vat_yuan"),
              49: ("增值税附加", "vat_surcharge_yuan"), 50: ("印花税", "stamp_tax_yuan"),
              63: ("资本金现金流", "equity_cashflow_yuan"),
              76: ("项目现金流", "project_cashflow_yuan")}
    for item in result["yearly"]:
        column = get_column_letter(int(item["year"]) + 4)
        for row, (label, key) in fields.items():
            unit = "比例" if key == "eol" else "万元"
            add(f"第{item['year']}年 {label}", "财务指标", f"{column}{row}",
                item[key] if unit == "比例" else item[key] / 10000, unit)
    return cases


def write_native_review(workbook, result: dict, edits: list[tuple]) -> None:
    for title in ("原版对账", "输入映射记录"):
        if title in workbook.sheetnames:
            del workbook[title]
    sheet = workbook.create_sheet("原版对账")
    sheet.append(["原版模板与服务器结果对账"])
    sheet.merge_cells("A1:G1")
    sheet.append(["已知规则差异尚未验收。以下公式需 Excel/WPS 重算，空缓存不表示通过。"])
    sheet.merge_cells("A2:G2")
    for location, detail in NATIVE_DIFFERENCES:
        sheet.append([location, detail])
        sheet.merge_cells(start_row=sheet.max_row, start_column=2, end_row=sheet.max_row, end_column=7)
    header = sheet.max_row + 2
    for column, text in enumerate(["指标", "模板位置", "模板计算值", "服务器快照", "差额", "检查", "单位"], 1):
        sheet.cell(header, column, text)
    for row, case in enumerate(comparison_cases(result), header + 1):
        ref = f"'{case['sheet']}'!{case['cell']}"
        values = [case["label"], ref, f"={ref}", case["expected"],
                  f'=IF(ISNUMBER(C{row}),C{row}-D{row},NA())',
                  f'=IF(ISNUMBER(C{row}),IF(ABS(E{row})<={case["tolerance"]},"一致","有差异"),"待复核")',
                  case["unit"]]
        if case["expected"] is None:
            values[3:6] = ["未定义", None, "服务器 IRR 未定义"]
        for column, value in enumerate(values, 1):
            sheet.cell(row, column, value)
        for column in (3, 4, 5):
            sheet.cell(row, column).number_format = "0.0000%" if case["unit"] == "比例" else "#,##0.0000"
    _style_review(sheet, header, [36, 30, 20, 20, 20, 22, 10])

    sheet = workbook.create_sheet("输入映射记录")
    sheet.append(["原版模板输入修改记录"])
    sheet.merge_cells("A1:D1")
    sheet.append(["年度金额覆盖原生中间收益，不反推未提供的电价或调频参数；费用仅按设备/其他两类归集。"])
    sheet.merge_cells("A2:D2")
    sheet.append(["工作表", "单元格", "模板原值/公式", "本次输入/公式"])
    for name, cell, old, new in edits:
        sheet.append([name, cell, str(old) if old is not None else "（空）", str(new)])
        # Display formula text without evaluating the original in this sheet.
        for column in (3, 4):
            sheet.cell(sheet.max_row, column).data_type = "s"
    _style_review(sheet, 3, [24, 12, 60, 52])


def _style_review(sheet, header: int, widths: list[int]):
    sheet.sheet_view.showGridLines = False
    sheet.freeze_panes = f"C{header + 1}"
    for cells in sheet:
        sheet.row_dimensions[cells[0].row].height = 34
        for cell in cells:
            cell.font = Font(name="Microsoft YaHei", size=11, color="173B35")
            cell.alignment = Alignment(vertical="center", wrap_text=True)
    for cell in sheet[header]:
        cell.fill = PatternFill("solid", fgColor="0B6B53")
        cell.font = Font(name="Microsoft YaHei", color="FFFFFF", bold=True)
    sheet["A1"].font = Font(name="Microsoft YaHei", size=16, bold=True, color="173B35")
    sheet.row_dimensions[2].height = 44
    for column, width in enumerate(widths, 1):
        sheet.column_dimensions[get_column_letter(column)].width = width
    sheet.auto_filter.ref = f"A{header}:{get_column_letter(len(widths))}{sheet.max_row}"
    sheet.print_title_rows = f"{header}:{header}"
    sheet.page_setup.orientation = "landscape"
    sheet.page_setup.paperSize = sheet.PAPERSIZE_A3
    sheet.page_setup.fitToWidth = 1
    sheet.page_setup.fitToHeight = 0
    sheet.sheet_properties.pageSetUpPr.fitToPage = True
    sheet.print_area = sheet.dimensions


def _formula_value(value):
    if hasattr(value, "ref"):
        return {"type": type(value).__name__, **dict(value), "text": getattr(value, "text", None)}
    return value


def _same_preserved_value(left, right) -> bool:
    """Compare persisted input values without flagging Excel's float rounding."""
    if isinstance(left, (int, float)) and not isinstance(left, bool) and \
            isinstance(right, (int, float)) and not isinstance(right, bool):
        return isclose(float(left), float(right), rel_tol=1e-10, abs_tol=1e-9)
    return left == right


def native_preservation_report(source: str | Path, exported: str | Path,
                               allowed_changes: dict[str, dict]) -> dict:
    """Compare persisted contents, including VBA bytes, rather than archive presence."""
    original = load_workbook(source)
    output = load_workbook(exported)
    changes, unexpected, styles, structures = [], [], [], []
    formula_count = 0
    try:
        for sheet in original:
            if sheet.title not in output:
                structures.append(f"missing sheet: {sheet.title}")
                continue
            target = output[sheet.title]
            if str(sheet.merged_cells) != str(target.merged_cells):
                structures.append(f"merges: {sheet.title}")
            for cells in sheet:
                for cell in cells:
                    other = target[cell.coordinate]
                    formula_count += cell.data_type == "f"
                    old, new = _formula_value(cell.value), _formula_value(other.value)
                    if old != new:
                        entry = {"sheet": sheet.title, "cell": cell.coordinate, "before": old, "after": new}
                        changes.append(entry)
                        if (cell.coordinate not in allowed_changes.get(sheet.title, {})
                                or not _same_preserved_value(
                                    new, allowed_changes[sheet.title][cell.coordinate])):
                            unexpected.append(entry)
                    if any(copy(getattr(cell, key)) != copy(getattr(other, key)) for key in
                           ("font", "fill", "border", "alignment", "number_format", "protection")):
                        styles.append(f"{sheet.title}!{cell.coordinate}")
        with ZipFile(source) as left, ZipFile(exported) as right:
            macros = [name for name in left.namelist() if name.lower().endswith("vbaproject.bin")]
            vba_equal = all(name in right.namelist() and left.read(name) == right.read(name) for name in macros)
            vba = "PRESERVED" if macros and vba_equal else "MISSING_OR_CHANGED" if macros else "ABSENT_IN_SOURCE"
        return {"source_sha256": sha256(Path(source).read_bytes()).hexdigest(),
                "export_sha256": sha256(Path(exported).read_bytes()).hexdigest(),
                "vba_status": vba, "source_formula_cells": formula_count,
                "changes": changes, "unexpected_changes": unexpected,
                "style_changes": styles, "structure_changes": structures,
                "status": "PRESERVED_WITH_DECLARED_CHANGES" if not (unexpected or styles or structures)
                and vba != "MISSING_OR_CHANGED" else "CHECK"}
    finally:
        original.close()
        output.close()


def review_native_cached_values(path: str | Path, result: dict,
                                *, engine_executed: bool = False,
                                native_engine: dict | None = None) -> dict:
    """Read saved caches, optionally after an external native recalculation."""
    workbook = load_workbook(path, data_only=True)
    rows = []
    try:
        identity = workbook["Banboos2.0快照"]["B5"].value if "Banboos2.0快照" in workbook else None
        for case in comparison_cases(result):
            value = workbook[case["sheet"]][case["cell"]].value if case["sheet"] in workbook else None
            expected = case["expected"]
            numeric = isinstance(value, (int, float)) and not isinstance(value, bool) and isfinite(value)
            delta = value - expected if numeric and expected is not None else None
            status = ("PENDING" if value is None else "UNDEFINED" if expected is None
                      else "ERROR" if not numeric else "MATCH" if abs(delta) <= case["tolerance"]
                      else "DIFFERENCE")
            rows.append(case | {"cached_value": value, "difference": delta, "status": status,
                                "difference_group": _difference_group(case["sheet"], case["cell"])})
        counts = {status: sum(row["status"] == status for row in rows)
                  for status in ("MATCH", "DIFFERENCE", "PENDING", "ERROR", "UNDEFINED")}
        difference_groups = {}
        for row in rows:
            if row["status"] == "DIFFERENCE":
                group = row["difference_group"]
                difference_groups[group] = difference_groups.get(group, 0) + 1
        matches_run = bool(result.get("run_id")) and identity == result["run_id"]
        status = ("RUN_MISMATCH" if not matches_run else "DIFFERENCES" if counts["DIFFERENCE"] or counts["ERROR"]
                  else "PENDING_RECALCULATION" if counts["PENDING"] or counts["UNDEFINED"]
                  else "CACHED_VALUES_MATCH")
        native_mode = (result.get("input_parameters") or {}).get("eol_method") == "native_xlsm"
        return {"run_id": result.get("run_id"), "workbook_run_id": identity, "status": status,
                "engine_executed": engine_executed,
                "native_recalculation_engine": native_engine,
                "formula_equivalence_verified": False,
                "known_rule_differences": [] if native_mode else NATIVE_DIFFERENCES,
                "compatibility_mode": "native_xlsm" if native_mode else "desktop_template_or_other",
                "counts": counts,
                "difference_groups": difference_groups, "checks": rows}
    finally:
        workbook.close()
