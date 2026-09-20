"""Prepare an Excel/WPS recalculation checklist for exported financial workbooks."""
from __future__ import annotations

from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

THIN = Side(style="thin", color="C9D8D2")
TEXT = "173B35"
GREEN = "0B6B53"
PASS = "E2F0D9"
WARN = "FFF2CC"


def write_recalc_audit(workbook, result: dict) -> None:
    """Add formula-vs-Python checks that become evaluable after Excel/WPS recalculates."""
    sheet = workbook.create_sheet("重算复核")
    sheet.cell(1, 1, "Excel/WPS 重算复核")
    sheet.cell(1, 1).font = Font(name="Microsoft YaHei", size=16, bold=True, color=TEXT)
    sheet.merge_cells("A1:F1")
    sheet.cell(2, 1, "打开并保存一次后，D 列显示公式与 Python 期望值的差额，E 列自动给出 PASS/CHECK。openpyxl 不执行 Excel 公式。")
    sheet.merge_cells("A2:F2")
    headers = ["检查项", "模板公式结果", "Python 期望值", "重算后差额", "状态", "说明"]
    for column, value in enumerate(headers, 1):
        cell = sheet.cell(4, column, value)
        cell.fill = PatternFill("solid", fgColor=GREEN)
        cell.font = Font(name="Microsoft YaHei", bold=True, color="FFFFFF")
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = Border(bottom=THIN)

    yearly = result.get("yearly", [])
    expected = [
        ("全投资 IRR", "='财务模板'!D78", result.get("full_irr"), "比例"),
        ("全投资 NPV（万元）", "='财务模板'!D79", _yuan_to_wan(result.get("full_npv_yuan")), "万元"),
        ("资本金 IRR", "='财务模板'!D65", result.get("equity_irr"), "比例"),
        ("资本金 NPV（万元）", "='财务模板'!D66", _yuan_to_wan(result.get("equity_npv_yuan")), "万元"),
    ]
    if yearly:
        expected.extend([
            ("首年收益合计（万元）", "='财务模板'!E27", _yuan_to_wan(yearly[0].get("revenue_yuan")), "万元"),
            ("首年实际应缴增值税（万元）", "='财务模板'!E48", _yuan_to_wan(yearly[0].get("actual_vat_yuan")), "万元"),
            ("首年项目现金流（万元）", "='财务模板'!E76", _yuan_to_wan(yearly[0].get("project_cashflow_yuan")), "万元"),
            ("首年资本金现金流（万元）", "='财务模板'!E63", _yuan_to_wan(yearly[0].get("equity_cashflow_yuan")), "万元"),
        ])
    for row, (label, formula, value, unit) in enumerate(expected, 5):
        sheet.cell(row, 1, label)
        sheet.cell(row, 2, formula)
        sheet.cell(row, 3, value)
        sheet.cell(row, 4, f"=B{row}-C{row}")
        sheet.cell(row, 5, f'=IF(ABS(D{row})<=IF("{unit}"="比例",0.0000001,0.01),"PASS","CHECK")')
        sheet.cell(row, 6, "Excel/WPS 重算后检查；服务器导出阶段仅保留公式和期望值。")
        for column in range(1, 7):
            cell = sheet.cell(row, column)
            cell.font = Font(name="Microsoft YaHei", color=TEXT, bold=column == 5)
            cell.fill = PatternFill("solid", fgColor=PASS if column == 5 else "FFFFFF")
            cell.border = Border(bottom=THIN)
            cell.alignment = Alignment(vertical="center", wrap_text=True)
        for column in (2, 3, 4):
            sheet.cell(row, column).number_format = "0.0000%" if unit == "比例" else "#,##0.00;[Red](#,##0.00);0.00"

    base_row = 5 + len(expected) + 2
    checks = [
        ("自动重算设置", workbook.calculation.calcMode == "auto" and workbook.calculation.fullCalcOnLoad
         and workbook.calculation.forceFullCalc, "auto / fullCalcOnLoad / forceFullCalc"),
        ("外部链接", not getattr(workbook, "_external_links", []), "未发现外部链接"),
        ("宏与数据表", False, "XLSX 不携带 VBA 和 DataTableFormula，需在原始 xlsm 中单独验证"),
        ("数组公式", False, "服务器模板页使用可重算普通公式，原始回收期数组公式需 Excel/WPS 专项验证"),
    ]
    for row, (label, ok, note) in enumerate(checks, base_row):
        values = [label, "PASS" if ok else "WARN", note]
        for column, value in enumerate(values, 1):
            cell = sheet.cell(row, column, value)
            cell.font = Font(name="Microsoft YaHei", color=TEXT, bold=column == 2)
            cell.fill = PatternFill("solid", fgColor=PASS if ok else WARN)
            cell.border = Border(bottom=THIN)
            cell.alignment = Alignment(vertical="center", wrap_text=True)
        sheet.merge_cells(start_row=row, start_column=3, end_row=row, end_column=6)
    sheet.auto_filter.ref = f"A4:F{base_row + len(checks) - 1}"
    sheet.freeze_panes = "A5"
    for column, width in {"A": 26, "B": 22, "C": 22, "D": 22, "E": 12, "F": 58}.items():
        sheet.column_dimensions[column].width = width


def _yuan_to_wan(value) -> float | None:
    return None if value is None else value / 10000
