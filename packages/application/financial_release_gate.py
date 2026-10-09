"""Release gate for the 1.6.6-compatible financial XLSM export.

The gate deliberately separates checks that can be proven by Python from checks
that require an installed spreadsheet engine.  This keeps a release report
useful on CI and makes a missing WPS/LibreOffice installation visible instead
of silently treating an Excel-only check as cross-engine validation.
"""
from __future__ import annotations

from collections.abc import Mapping
from pathlib import Path
from zipfile import ZipFile

from openpyxl import load_workbook

from packages.application.financial_template_xlsm import template_sheet_values
from packages.application.financial_xlsm_review import (
    native_preservation_report,
    review_native_cached_values,
)
from packages.application.spreadsheet_engine import detect_spreadsheet_engines, recalculate_xlsm


def inspect_workbook_features(path: str | Path) -> dict[str, object]:
    """Inspect workbook features that can be lost by a non-native writer."""
    workbook_path = Path(path).expanduser().resolve()
    names: list[str]
    with ZipFile(workbook_path) as archive:
        names = archive.namelist()
        xml = b"\n".join(
            archive.read(name) for name in names if name.endswith(".xml")
        ).decode("utf-8", errors="ignore")
    workbook = load_workbook(workbook_path, data_only=False, keep_vba=True, read_only=False)
    try:
        formula_count = 0
        array_formula_count = 0
        for sheet in workbook.worksheets:
            for row in sheet.iter_rows():
                for cell in row:
                    if cell.data_type == "f":
                        formula_count += 1
                    if type(cell.value).__name__ == "ArrayFormula":
                        array_formula_count += 1
        defined_names = len(workbook.defined_names)
        sheet_names = list(workbook.sheetnames)
    finally:
        workbook.close()
    vba = [name for name in names if name.lower().endswith("vbaproject.bin")]
    external_links = [name for name in names if name.startswith("xl/externalLinks/")]
    tables = [name for name in names if name.startswith("xl/tables/") and name.endswith(".xml")]
    data_table_formulas = xml.count("<dataTable") + xml.upper().count("TABLE(")
    connections = [name for name in names if name == "xl/connections.xml"]
    return {
        "path": str(workbook_path),
        "sheets": sheet_names,
        "formula_count": formula_count,
        "array_formula_count": array_formula_count,
        "defined_name_count": defined_names,
        "vba_project_count": len(vba),
        "excel_table_count": len(tables),
        "external_link_count": len(external_links),
        "connection_count": len(connections),
        "data_table_formula_count": data_table_formulas,
        "vba_status": "present" if vba else "absent",
        "data_table_status": "present" if data_table_formulas else "absent",
        "external_link_status": "present" if external_links else "absent",
    }


def _check(status: str, detail: str, **extra: object) -> dict[str, object]:
    return {"status": status, "detail": detail, **extra}


def run_release_gate(
    source_template: str | Path,
    exported_workbook: str | Path,
    result: Mapping[str, object],
    *,
    recalculate: bool = True,
    recalc_output: str | Path | None = None,
) -> dict[str, object]:
    """Run preservation, native cache and cross-engine readiness checks."""
    source = Path(source_template).expanduser().resolve()
    exported = Path(exported_workbook).expanduser().resolve()
    allowed = template_sheet_values(dict(result))
    preservation = native_preservation_report(source, exported, allowed)
    features = inspect_workbook_features(exported)
    engines = detect_spreadsheet_engines()
    if recalc_output is None:
        recalc_path = exported.with_name(f"{exported.stem}.recalculated.xlsm")
    else:
        recalc_path = Path(recalc_output).expanduser().resolve()

    native = {"status": "skipped", "engine": engines}
    reviewed_path = exported
    if recalculate:
        native = recalculate_xlsm(exported, recalc_path)
        if native.get("status") == "recalculated":
            reviewed_path = recalc_path
    reconciliation = review_native_cached_values(
        reviewed_path,
        dict(result),
        engine_executed=native.get("status") == "recalculated",
        native_engine=native.get("engine") if native.get("status") == "recalculated" else None,
    )

    preservation_ok = preservation["status"] == "PRESERVED_WITH_DECLARED_CHANGES"
    native_ok = native.get("status") == "recalculated"
    reconciliation_ok = reconciliation["status"] == "CACHED_VALUES_MATCH"
    alternate_engines = {
        name: path for name, path in (engines.get("engines") or {}).items()
        if name in {"wps", "libreoffice"}
    }
    checks = {
        "template_preservation": _check(
            "PASS" if preservation_ok else "FAIL",
            "原始工作表、样式、结构与 VBA 状态已检查",
            report=preservation,
        ),
        "native_recalculation": _check(
            "PASS" if native_ok else "PENDING" if native.get("status") == "unavailable" else "FAIL",
            str(native.get("status")),
            report=native,
        ),
        "native_value_reconciliation": _check(
            "PASS" if reconciliation_ok else "PENDING" if reconciliation["status"] == "PENDING_RECALCULATION" else "FAIL",
            str(reconciliation["status"]),
            report=reconciliation,
        ),
        "workbook_feature_inspection": _check(
            "PASS", "已检查宏、数据表、数组公式、外部链接和连接",
            report=features,
        ),
        "cross_engine_availability": _check(
            "PASS" if alternate_engines else "PENDING",
            "已发现 WPS/LibreOffice" if alternate_engines else "当前主机未发现 WPS 或 LibreOffice",
            available=alternate_engines,
        ),
    }
    if all((preservation_ok, native_ok, reconciliation_ok, bool(alternate_engines))):
        status = "RELEASE_READY"
    elif preservation_ok and native_ok and reconciliation_ok and not alternate_engines:
        status = "PENDING_CROSS_ENGINE"
    else:
        status = "BLOCKED"
    return {
        "status": status,
        "source_template": str(source),
        "exported_workbook": str(exported),
        "recalculated_workbook": str(reviewed_path) if reviewed_path != exported else None,
        "checks": checks,
        "engines": engines,
    }
